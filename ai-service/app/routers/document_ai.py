"""文档 AI 路由。

POST /ai/documents/classify      - 自动分类 / 打标签 / 摘要
POST /ai/documents/generate-pdf  - 生成报价 PDF
POST /ai/documents/generate-docx - 生成技术协议 DOCX
"""
from __future__ import annotations

import logging
import os
import re
import asyncio
from datetime import datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Security, HTTPException
from pydantic import BaseModel, Field

from app.core.auth import verify_api_key
from app.core.config import settings
from app.core.json_utils import extract_json
from app.core.llm_router import llm_router

logger = logging.getLogger(__name__)
router = APIRouter(dependencies=[Security(verify_api_key)])

EXPORT_DIR = os.path.join(settings.data_dir, "exports")

_VALID_CATEGORIES = ["产品知识", "技术协议", "操作手册", "FAQ", "内部制度"]


class _DepMissing(Exception):
    """可选依赖（reportlab / python-docx 等）未安装时抛出，用于在线程池内传递。"""


class ClassifyRequest(BaseModel):
    text: str = Field(..., max_length=50_000)


class QuotationItem(BaseModel):
    name: str = Field(..., max_length=500)
    quantity: float = Field(default=1, ge=0)
    unitPrice: float = Field(default=0, ge=0)
    totalPrice: float = Field(default=0, ge=0)


class GeneratePdfRequest(BaseModel):
    title: str = Field(default="报价单", max_length=200)
    customer: str = Field(default="", max_length=200)
    items: List[QuotationItem] = Field(default_factory=list)
    totalAmount: float = Field(default=0, ge=0)
    validUntil: Optional[str] = Field(default=None, max_length=50)
    terms: Optional[str] = Field(default=None, max_length=5_000)
    companyName: str = Field(default="", max_length=200)


class KvPair(BaseModel):
    key: str = Field(..., max_length=200)
    value: str = Field(..., max_length=5_000)


class GenerateDocxRequest(BaseModel):
    title: str = Field(default="技术协议", max_length=200)
    customer: str = Field(default="", max_length=200)
    solution: Optional[str] = Field(default=None, max_length=5_000)
    params: List[KvPair] = Field(default_factory=list)
    config: List[KvPair] = Field(default_factory=list)
    acceptance: Optional[str] = Field(default=None, max_length=5_000)


@router.post("/classify")
async def classify_document(request: ClassifyRequest):
    """自动分类、打标签并生成摘要。"""
    text = request.text.strip()
    if not text:
        return {"category": "产品知识", "tags": [], "summary": ""}

    system_prompt = (
        "你是文档分析助手。请对给定文本进行分类、提取标签并生成摘要。"
        "类别只能从以下选项中选择：产品知识、技术协议、操作手册、FAQ、内部制度。"
        "输出 JSON: "
        "{\"category\": '类别', \"tags\": ['标签1','标签2'], \"summary\": '200字以内摘要'}。"
    )
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": text[:8_000]},
    ]

    try:
        response = await llm_router.chat(
            messages=messages,
            module="document",
            temperature=0.2,
            response_format={"type": "json_object"},
        )
        data = extract_json(response)
        category = data.get("category") or "产品知识"
        if category not in _VALID_CATEGORIES:
            category = "产品知识"
        tags = data.get("tags") or []
        if not isinstance(tags, list):
            tags = [str(tags)]
        summary = (data.get("summary") or "").strip()
        return {"category": category, "tags": tags, "summary": summary}
    except Exception:
        logger.exception("classify_document LLM call failed")
        return {"category": "产品知识", "tags": [], "summary": ""}


def _safe_filename(name: str) -> str:
    return re.sub(r"[\\/:*?\"<>|]+", "_", name).strip("_") or "export"


@router.post("/generate-pdf")
async def generate_pdf(request: GeneratePdfRequest):
    """使用 reportlab 生成报价单 PDF（懒加载）。"""
    os.makedirs(EXPORT_DIR, exist_ok=True)

    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    safe_title = _safe_filename(request.title)
    file_name = f"quotation_{safe_title}_{stamp}.pdf"
    file_path = os.path.join(EXPORT_DIR, file_name)

    def _build() -> str:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError:
            logger.error("reportlab is not installed")
            raise _DepMissing("PDF 生成依赖 reportlab 未安装")

        doc = SimpleDocTemplate(
            file_path, pagesize=A4,
            leftMargin=20 * mm, rightMargin=20 * mm,
            topMargin=18 * mm, bottomMargin=18 * mm,
        )
        styles = getSampleStyleSheet()
        title_style = styles["Title"]
        normal_style = styles["Normal"]
        cell_style = ParagraphStyle(
            "cell", parent=normal_style, fontSize=9, leading=12
        )

        elements: List[Any] = []
        elements.append(Paragraph(request.title or "报价单", title_style))
        elements.append(Spacer(1, 6 * mm))

        meta_lines = []
        if request.companyName:
            meta_lines.append(f"报价单位：{request.companyName}")
        if request.customer:
            meta_lines.append(f"客户名称：{request.customer}")
        meta_lines.append(f"报价日期：{datetime.now().strftime('%Y-%m-%d')}")
        if request.validUntil:
            meta_lines.append(f"有效期至：{request.validUntil}")
        for line in meta_lines:
            elements.append(Paragraph(line, normal_style))
        elements.append(Spacer(1, 4 * mm))

        header = ["序号", "名称", "数量", "单价(元)", "小计(元)"]
        table_data = [header]
        for idx, item in enumerate(request.items, start=1):
            table_data.append([
                str(idx),
                Paragraph(str(item.name), cell_style),
                f"{item.quantity:g}",
                f"{item.unitPrice:.2f}",
                f"{item.totalPrice:.2f}",
            ])
        total = request.totalAmount or sum(i.totalPrice for i in request.items)
        table_data.append(["", Paragraph("合计", cell_style), "", "", f"{total:.2f}"])

        col_widths = [14 * mm, 70 * mm, 22 * mm, 32 * mm, 32 * mm]
        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E5AAC")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.whitesmoke, colors.white]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8EEF7")),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 6 * mm))

        if request.terms:
            elements.append(Paragraph("<b>条款说明</b>", normal_style))
            elements.append(Paragraph(request.terms.replace("\n", "<br/>"), cell_style))

        doc.build(elements)
        return file_name

    try:
        # reportlab 构建 PDF 是 CPU 密集同步操作，丢进线程池
        await asyncio.to_thread(_build)
        return {"url": f"/exports/{file_name}", "fileName": file_name}
    except _DepMissing as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception:
        logger.exception("generate_pdf failed")
        raise HTTPException(status_code=500, detail="PDF 生成失败")


@router.post("/generate-docx")
async def generate_docx(request: GenerateDocxRequest):
    """使用 python-docx 生成技术协议 DOCX（懒加载）。"""
    os.makedirs(EXPORT_DIR, exist_ok=True)

    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    safe_title = _safe_filename(request.title)
    file_name = f"agreement_{safe_title}_{stamp}.docx"
    file_path = os.path.join(EXPORT_DIR, file_name)

    def _build() -> str:
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            logger.error("python-docx is not installed")
            raise _DepMissing("DOCX 生成依赖 python-docx 未安装")

        document = Document()

        # 标题
        h = document.add_heading(request.title or "技术协议", level=0)
        for run in h.runs:
            run.font.size = Pt(20)

        # 基础信息
        info_table = document.add_table(rows=0, cols=2)
        info_table.style = "Light Grid Accent 1"
        info_rows: List[tuple] = []
        if request.customer:
            info_rows.append(("甲方（客户）", request.customer))
        info_rows.append(("签订日期", datetime.now().strftime('%Y-%m-%d')))
        for k, v in info_rows:
            row = info_table.add_row().cells
            row[0].text = k
            row[1].text = str(v)

        document.add_paragraph("")

        # 方案概述
        if request.solution:
            document.add_heading("一、方案概述", level=1)
            document.add_paragraph(request.solution)

        # 技术参数
        def _write_kv_table(heading: str, pairs: List[KvPair]):
            if not pairs:
                return
            document.add_heading(heading, level=1)
            tbl = document.add_table(rows=1, cols=2)
            tbl.style = "Light Grid Accent 1"
            hdr = tbl.rows[0].cells
            hdr[0].text = "项目"
            hdr[1].text = "内容"
            for p in pairs:
                cells = tbl.add_row().cells
                cells[0].text = p.key
                cells[1].text = p.value

        _write_kv_table("二、技术参数", request.params)
        _write_kv_table("三、配置清单", request.config)

        # 验收标准
        if request.acceptance:
            document.add_heading("四、验收标准", level=1)
            document.add_paragraph(request.acceptance)

        document.add_heading("五、其他约定", level=1)
        document.add_paragraph(
            "本协议一式两份，甲乙双方各执一份，自双方签字盖章之日起生效。"
            "未尽事宜由双方协商解决。"
        )

        document.add_paragraph("")
        sign_table = document.add_table(rows=2, cols=2)
        sign_table.style = "Light List Accent 1"
        sign_table.rows[0].cells[0].text = "甲方（盖章）："
        sign_table.rows[0].cells[1].text = "乙方（盖章）："
        sign_table.rows[1].cells[0].text = "日期："
        sign_table.rows[1].cells[1].text = "日期："

        document.save(file_path)
        return file_name

    try:
        # python-docx 生成是同步操作，丢进线程池
        await asyncio.to_thread(_build)
        return {"url": f"/exports/{file_name}", "fileName": file_name}
    except _DepMissing as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception:
        logger.exception("generate_docx failed")
        raise HTTPException(status_code=500, detail="DOCX 生成失败")
