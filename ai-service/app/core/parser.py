import io
import json
import logging
import re
from typing import List

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from markdown import markdown
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# ── 文本可读性判定 ─────────────────────────────────────────────
#
# 为什么需要这个：原来 parse_document 的 else 分支对**任何**未知类型都做
# `content.decode("utf-8", errors="ignore")`。二进制文件（.doc 是 OLE 复合文档、
# .wps、.zip…）会被硬解成一堆内部标记，例如
#     "hN h$ B*CJ H*OJ PJ QJ aJ o(ph"
# 这些垃圾照样分块入库，本机实测**一个 .doc 产出 1109 块、其中 1108 块是垃圾，
# 占整个向量库的 44%**。检索时它们和真资料抢 Top-K，回答自然不准。
#
# 边界要设在入口：解析不出人能读的文本，就明确失败，而不是"尽量返回点什么"。

_CJK = re.compile(r"[一-鿿]")
_WORD = re.compile(r"[A-Za-z]{3,}")
# Word/RTF/PDF 抽取失败时的典型残留
_BINARY_MARK = re.compile(r"OJ\s*PJ|B\*CJ|nonmarkingreturn|/\.null|\\par\b|CJ\s+OJ")


def looks_like_text(text: str, min_length: int = 20) -> bool:
    """判断一段抽取结果是否是人能读的文本。

    两条独立证据满足其一即可：
      · 中文密度 >= 15%（正常中文技术文档远高于此）
      · 像正常英文：ASCII 字母 >= 60% 且至少 5 个长度 >=3 的单词

    在本机 531 块业务对象内容上跑过，误杀 0 条。
    """
    t = (text or "").strip()
    if len(t) < min_length:
        return False
    if _BINARY_MARK.search(t):
        return False

    cjk = len(_CJK.findall(t))
    if cjk / len(t) >= 0.15:
        return True

    letters = sum(1 for c in t if c.isascii() and c.isalpha())
    if letters / len(t) >= 0.60 and len(_WORD.findall(t)) >= 5:
        return True
    return False


# 解析失败时后端回退写入的占位文案（documents.service.fallbackParse），
# 原文形如 `[图纸A-01.dwg] 文件已上传，等待解析。`
#
# 必须匹配完整形态，不能只认"等待解析"四个字：正文里完全可能出现
# "有五份因为是扫描件而处于等待解析状态"这样的正常表述，
# 只认关键词会把真资料一起误杀（这条是被测试用例抓出来的）。
_PLACEHOLDER = re.compile(r"文件已上传[，,]?\s*等待解析")


def is_parse_placeholder(text: str) -> bool:
    """识别"文件已上传，等待解析"这类占位文案。

    后端解析失败时会回退成 `[xxx.dwg] 文件已上传，等待解析。`，
    这段话中文密度很高，能骗过 looks_like_text，于是被当成正文分块入库。
    后果是检索里凭空多出一批"等待解析"的噪声块 —— 用户问什么都可能命中它们，
    而它们不含任何业务信息。
    """
    t = (text or "").strip()
    # 占位文案就一句话；正文里引用到这句话的情况通常篇幅长得多
    if len(t) > 120:
        return False
    return bool(_PLACEHOLDER.search(t))


def is_index_status_marker(text: str) -> bool:
    """识别"索引状态标记"被当成正文的情况。

    知识资料的 parsedText 存的不是正文，而是一段状态 JSON：
        {"indexed":false,"chunks":0,"reason":"scanned_pdf","note":"该文件为扫描件…"}
    历史上有一次全量重建把它当正文索引了（见 reindex.service.ts 的说明），
    本机至今还留着 5 条这样的残渣。

    这种内容骗得过 looks_like_text —— note 字段里有中文，中文密度足够高 ——
    所以要单独认一下签名。
    """
    t = (text or "").strip()
    if '"indexed"' not in t and '"indexedAt"' not in t:
        return False
    start = t.find("{")
    if start < 0:
        return False
    try:
        obj = json.loads(t[start:])
    except (ValueError, TypeError):
        return False
    return isinstance(obj, dict) and ("indexed" in obj or "indexedAt" in obj)


def is_char_per_line(text: str) -> bool:
    """逐字符换行 —— PDF 抽取失败的另一种典型形态（"5\\n/\\n1\\n3\\n."）。

    这种内容字面上"可读"，但已经失去所有词与句的结构，
    向量化后是纯噪声，检索里只会添乱。
    """
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if len(lines) < 8:
        return False
    return sum(1 for ln in lines if len(ln) <= 2) / len(lines) > 0.6


# 中文文档常见编码。utf-8 失败时按序尝试，取第一个"读起来正常"的结果，
# 而不是 errors="ignore" 硬吞 —— 后者会把 GBK 文本变成一片乱码。
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "utf-16", "big5")


def decode_text(content: bytes, filename: str = "") -> str:
    """按多种编码尝试解码纯文本文件，挑可读性最好的。"""
    best_text, best_score = "", -1.0
    for enc in _TEXT_ENCODINGS:
        try:
            text = content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        if looks_like_text(text):
            return text
        # 都不理想时，用"可识别字符占比"挑一个相对最好的
        score = sum(1 for c in text if c.isprintable() or c.isspace()) / max(len(text), 1)
        if score > best_score:
            best_text, best_score = text, score
    if not best_text:
        raise ValueError(f"无法识别文件编码：{filename or '未知文件'}")
    return best_text


# 明确支持的扩展名。不在这里的一律拒绝，不再"尽量解一下"。
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".xls", ".md", ".txt", ".csv")


def parse_document(filename: str, content: bytes) -> str:
    try:
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            text = parse_pdf(content)
        elif lower_name.endswith(".docx"):
            text = parse_docx(content)
        elif lower_name.endswith(".xlsx"):
            text = parse_excel(content)
        elif lower_name.endswith(".xls"):
            text = parse_xls(content)
        elif lower_name.endswith(".md"):
            text = parse_markdown(decode_text(content, filename))
        elif lower_name.endswith((".txt", ".csv")):
            text = decode_text(content, filename)
        elif lower_name.endswith(".doc"):
            # .doc 是 Word 97 的 OLE 二进制格式，python-docx 读不了，
            # 硬解字节只会得到内部标记。给一句用户能照做的话，别默默入库垃圾。
            raise ValueError(
                f"{filename} 是旧版 .doc 格式，无法直接解析。"
                "请用 Word 打开后「另存为」.docx 或 PDF 再上传。"
            )
        else:
            ext = lower_name.rsplit(".", 1)[-1] if "." in lower_name else "无扩展名"
            raise ValueError(
                f"不支持的文件类型（.{ext}）。"
                f"当前支持：{'、'.join(SUPPORTED_EXTENSIONS)}"
            )

        # 出口统一校验：解析器返回了内容，不代表内容是人能读的。
        # 扫描件 PDF、CID 字体 PDF 都会走到这里。
        if not looks_like_text(text):
            raise ValueError(
                f"{filename} 未能提取出可读文本（可能是扫描件、图片型 PDF 或加密文档）。"
                "请上传文字版文件。"
            )
        return text
    except ValueError:
        raise
    except Exception as exc:
        logger.error(f"Failed to parse document {filename}: {exc}")
        raise ValueError(f"Failed to parse document {filename}: {exc}") from exc


def parse_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        texts = []
        for page in reader.pages:
            texts.append(page.extract_text() or "")
        full_text = "\n".join(texts).strip()

        if len(full_text) >= 20:
            return full_text

        logger.info("PDF text extraction returned < 20 chars, attempting OCR fallback...")
        return _pdf_ocr_fallback(content)
    except Exception as exc:
        logger.error(f"Failed to parse PDF: {exc}")
        raise ValueError(f"Failed to parse PDF: {exc}") from exc


def _pdf_ocr_fallback(content: bytes) -> str:
    """扫描件 PDF OCR 回退：将每页转为图片后调用 PaddleOCR。

    recognize() 只接受**图片字节**，且内部自己落临时文件、自己懒加载引擎
    （见 media/ocr.py）—— 这里不要再自己建临时文件，也不要预热引擎，
    否则要么类型对不上，要么多写一份没人读的临时文件。
    """
    try:
        from app.core.media.ocr import recognize
        import fitz  # PyMuPDF
    except ImportError as exc:
        logger.warning(f"OCR fallback unavailable (missing dependency): {exc}")
        return ""

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        logger.warning(f"OCR fallback failed to open PDF: {exc}")
        return ""

    texts = []
    try:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")
            try:
                result = recognize(img_bytes, file_name=f"page_{i}.png")
            except Exception as exc:
                # 单页失败（引擎未装、该页图异常）不该让整份文件颗粒无收，
                # 记下来继续下一页。
                logger.warning(f"OCR fallback failed on page {i}: {exc}")
                continue
            if result and result.fullText:
                texts.append(result.fullText)
    finally:
        doc.close()

    full_text = "\n".join(texts).strip()
    if full_text:
        logger.info(f"OCR fallback extracted {len(full_text)} chars from {len(texts)} pages")
    else:
        logger.warning("OCR fallback also returned empty text")
    return full_text


def parse_docx(content: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(content))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            parts.append(_docx_table_to_text(table))
        return "\n".join(parts)
    except Exception as exc:
        logger.error(f"Failed to parse DOCX: {exc}")
        raise ValueError(f"Failed to parse DOCX: {exc}") from exc


def _docx_table_to_text(table) -> str:
    """将 Word 表格转为 Markdown 格式"""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def parse_excel(content: bytes) -> str:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        return _openpyxl_to_text(wb)
    except Exception as exc:
        logger.error(f"Failed to parse XLSX: {exc}")
        raise ValueError(f"Failed to parse XLSX: {exc}") from exc


def parse_xls(content: bytes) -> str:
    try:
        import xlrd

        wb = xlrd.open_workbook(file_contents=content)
        texts = []
        for sheet in wb.sheets():
            texts.append(f"Sheet: {sheet.name}")
            headers = None
            for row_idx in range(sheet.nrows):
                row = sheet.row(row_idx)
                values = [str(cell.value) for cell in row if cell.value is not None]
                if row_idx == 0:
                    headers = values
                    texts.append(" | ".join(values))
                    texts.append(" | ".join(["---"] * len(values)))
                else:
                    texts.append(" | ".join(values))
        return "\n".join(texts)
    except ImportError as exc:
        logger.error("xlrd is required for .xls support")
        raise ValueError("xlrd is required for .xls support") from exc
    except Exception as exc:
        logger.error(f"Failed to parse XLS: {exc}")
        raise ValueError(f"Failed to parse XLS: {exc}") from exc


def _openpyxl_to_text(wb) -> str:
    """Excel 转 Markdown 表格（保留表头 + 数据行）"""
    texts = []
    for sheet in wb.worksheets:
        texts.append(f"## Sheet: {sheet.title}")
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(c) if c is not None else "" for c in rows[0]]
        texts.append("| " + " | ".join(headers) + " |")
        texts.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            texts.append("| " + " | ".join(cells) + " |")
    return "\n".join(texts)


def parse_markdown(content: str) -> str:
    try:
        html = markdown(content)
        soup = BeautifulSoup(html, "html.parser")
        return soup.get_text()
    except Exception as exc:
        logger.error(f"Failed to parse Markdown: {exc}")
        raise ValueError(f"Failed to parse Markdown: {exc}") from exc


def chunk_text(text: str, chunk_size: int = 500, chunk_overlap: int = 100) -> List[str]:
    try:
        if not text:
            return []

        # 防御 overlap >= chunk_size 导致步长 <= 0 的死循环
        if chunk_overlap >= chunk_size:
            chunk_overlap = max(0, chunk_size // 4)

        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk = text[start:end]
            chunks.append(chunk.strip())
            start += chunk_size - chunk_overlap
            if end == text_len:
                break

        return [c for c in chunks if c]
    except Exception as exc:
        logger.error(f"Failed to chunk text: {exc}")
        raise ValueError(f"Failed to chunk text: {exc}") from exc
