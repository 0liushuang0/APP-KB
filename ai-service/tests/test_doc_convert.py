"""旧版 .doc: 本机有 LibreOffice 时自动转换后解析; 没有时给出可操作的提示"""
import io

import pytest
from docx import Document

from app.core import parser


def test_doc_without_converter_gives_actionable_message(monkeypatch):
    monkeypatch.setattr(parser, "convert_doc_to_docx", lambda content, timeout=90: None)
    with pytest.raises(ValueError) as e:
        parser.parse_document("协议.doc", b"\xd0\xcf\x11\xe0 binary")
    assert "另存为" in str(e.value) and "LibreOffice" in str(e.value)


def test_doc_converted_then_parsed(monkeypatch):
    d = Document()
    d.add_paragraph("电镀锡线1450组合拉矫机组技术规格书 Technical Specification of the tension leveller unit")
    d.add_paragraph("带钢厚度 0.15~0.6 mm，带钢宽度 700~1300 mm，Strip thickness and width are as follows.")
    buf = io.BytesIO(); d.save(buf)
    monkeypatch.setattr(parser, "convert_doc_to_docx", lambda content, timeout=90: buf.getvalue())
    text = parser.parse_document("规格书.doc", b"\xd0\xcf\x11\xe0 binary")
    assert "组合拉矫机组" in text


@pytest.mark.skipif(parser._find_soffice() is None, reason="本机没有 LibreOffice")
def test_real_libreoffice_roundtrip(tmp_path):
    """真实转换: 先用 LibreOffice 把 docx 转成 doc, 再走 .doc 解析路径"""
    import subprocess
    from pathlib import Path

    d = Document()
    d.add_paragraph("技术协议 Technical Agreement：卷取机最大卷重 30 吨，The maximum coil weight of the coiler is 30 tons.")
    d.add_paragraph("供货范围包括开卷机、卷取机、液压站和电气控制系统 including uncoiler, coiler and hydraulic station.")
    src = tmp_path / "a.docx"; d.save(src)
    subprocess.run([parser._find_soffice(), "--headless", "--convert-to", "doc", "--outdir", str(tmp_path), str(src)],
                   capture_output=True, timeout=120, check=True)
    doc_bytes = Path(tmp_path / "a.doc").read_bytes()
    text = parser.parse_document("a.doc", doc_bytes)
    assert "卷取机最大卷重" in text
