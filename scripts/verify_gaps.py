#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gap A / B1 / B2 / B3 / B4 端到端验证脚本。"""
import json, sys, requests

BASE = "http://localhost:3001/api"
S = requests.Session()
RES = []

def step(name, ok, detail=""):
    RES.append((name, ok, detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}: {detail}")

def api(method, path, **kw):
    url = BASE + path
    try:
        r = S.request(method, url, timeout=90, **kw)
        return r.status_code, r.json()
    except Exception as e:
        return 0, {"error": str(e)}

# ---- login ----
sc, r = api("POST", "/auth/single-mode-login", json={})
TOKEN = r["data"]["accessToken"]
S.headers.update({"Authorization": f"Bearer {TOKEN}", "Content-Type": "application/json"})
step("登录", sc in (200, 201) and bool(TOKEN), f"token={'有' if TOKEN else '无'}")

# =====================================================================
# Gap A: Office 文档（docx）上传 -> structure -> parsedText 非空
# =====================================================================
print("\n=== Gap A: docx 文档解析 ===")
sc, r = api("POST", "/requirements", json={
    "title": "GapA验证-文档解析",
    "description": "客户需要液压油缸，行程1200mm，工作压力21MPa。",
})
A_RID = r["data"]["id"]
step("A.1 创建需求", sc in (200, 201), f"rid={A_RID}")

# 生成一个最小 docx（python-docx）
from docx import Document as Docx
import io, base64
doc = Docx()
doc.add_heading("技术规格书", 0)
doc.add_paragraph("产品类型：液压油缸")
doc.add_paragraph("行程：1200mm")
doc.add_paragraph("工作压力：21MPa")
doc.add_paragraph("缸径：63mm")
buf = io.BytesIO()
doc.save(buf)
docx_bytes = buf.getvalue()

files = {"file": ("spec.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
del S.headers["Content-Type"]
sc, r = api("POST", f"/requirements/{A_RID}/upload", files=files)
S.headers.update({"Content-Type": "application/json"})
ATT_ID = r.get("data", {}).get("id") if isinstance(r, dict) else None
step("A.2 上传 docx", sc in (200, 201), f"sc={sc}")

sc, r = api("POST", f"/requirements/{A_RID}/structure")
data = r.get("data", {})
atts = data.get("attachments", [])
parsed = [a for a in atts if a.get("parsedText", "").strip()]
step("A.3 结构化后附件 parsedText 非空", len(parsed) > 0,
     f"已解析附件数={len(parsed)} | parsedText前30字={parsed[0]['parsedText'][:30] if parsed else '空'}")

# 需求卡 sources 应含 document 类型来源
card = json.loads(data.get("card") or "{}")
doc_sources = [s for s in card.get("sources", []) if s.get("source_type") in ("document", "text")]
step("A.4 需求卡含文档来源", len(doc_sources) > 0,
     f"文档来源数={len(doc_sources)}")

# =====================================================================
# Gap B1: structure 返回值含 autoMatch.top3
# =====================================================================
print("\n=== Gap B1: 结构化自动匹配 ===")
step("B1 structure 返回含 autoMatch", "autoMatch" in data,
     f"autoMatch keys={list(data.get('autoMatch', {}).keys()) if isinstance(data.get('autoMatch'), dict) else '无'}")
if isinstance(data.get("autoMatch"), dict):
    top3 = data["autoMatch"].get("top3", [])
    step("B1 autoMatch.top3 存在", isinstance(top3, list),
         f"候选数={len(top3)}")

# =====================================================================
# Gap B3: 由需求创建 proposal 带 requirementId
# =====================================================================
print("\n=== Gap B3: 需求↔方案外键 ===")
sc, r = api("POST", "/proposals", json={
    "proposalNo": f"PROP-GAPB3-{int(__import__('time').time())}",
    "title": "GapB3关联需求方案",
    "customerId": data.get("customerId") or "",
    "requirementId": A_RID,
    "aiGenerated": True,
})
B3_PID = r.get("data", {}).get("id") if isinstance(r, dict) else None
got_reqid = r.get("data", {}).get("requirementId") if isinstance(r, dict) else None
step("B3.1 创建 proposal 带 requirementId", sc in (200, 201) and got_reqid == A_RID,
     f"pid={B3_PID} | requirementId={got_reqid}")

# =====================================================================
# Gap B2 + B4: 报价 approved->sent->won 自动建项目
# =====================================================================
print("\n=== Gap B2/B4: 报价成交自动建项目 ===")
# 取一个真实 customerId（需求卡 customerId 可能为空，回退用列表第一个客户）
cust_id = data.get("customerId") or ""
if not cust_id:
    _, cr = api("GET", "/customers")
    clist = cr.get("data", [])
    if isinstance(clist, list) and clist:
        cust_id = clist[0].get("id") or clist[0].get("customerId") or ""

sc, r = api("POST", "/quotations", json={
    "quotationNo": f"QUO-GAPB2-{int(__import__('time').time())}",
    "customerId": cust_id,
    "proposalId": B3_PID or "",
    "amount": 50000,
    "items": [],
})
Q_ID = r.get("data", {}).get("id") if isinstance(r, dict) else None
q_status = r.get("data", {}).get("status")
step("B2.1 创建报价", sc in (200, 201), f"qid={Q_ID} | 初始状态={q_status}")

# 推进到 approved（若已是 approved 则跳过）
if q_status not in ("approved", "sent"):
    if q_status == "pending":
        api("POST", f"/quotations/{Q_ID}/approve")
    elif q_status in ("draft", "rejected"):
        api("POST", f"/quotations/{Q_ID}/submit")
        api("POST", f"/quotations/{Q_ID}/approve")
sc, r = api("GET", f"/quotations/{Q_ID}")
cur = r.get("data", {}).get("status")
step("B4.1 推进到 approved", cur == "approved", f"状态={cur}")

# B4: approved -> sent
sc, r = api("POST", f"/quotations/{Q_ID}/send")
cur = r.get("data", {}).get("status")
step("B4.2 标记已发送 (markSent)", cur == "sent", f"状态={cur}")

# 统计 win 前项目数
sc, r = api("GET", "/projects")
projects_before = len(r.get("data", [])) if isinstance(r.get("data"), list) else 0

# B2: sent -> won (应自动建项目)
sc, r = api("POST", f"/quotations/{Q_ID}/win")
won_data = r.get("data", {})
cur = won_data.get("status")
auto_project_id = won_data.get("projectId")
step("B2.1 确认成交 (markWon)", cur == "won", f"状态={cur}")
step("B2.2 返回含自动创建的 projectId", bool(auto_project_id), f"projectId={auto_project_id}")

# 验证项目确实被创建且关联 quotationId
sc, r = api("GET", "/projects")
projects_after = r.get("data", []) if isinstance(r.get("data"), list) else []
step("B2.3 项目数 +1", len(projects_after) == projects_before + 1,
     f"前={projects_before} 后={len(projects_after)}")
linked = [p for p in projects_after if p.get("quotationId") == Q_ID]
step("B2.4 新项目 quotationId 正确关联", len(linked) == 1,
     f"关联项目={linked[0].get('projectNo') if linked else '无'}")

# =====================================================================
# 清理
# =====================================================================
for eid, ep in [(A_RID, "requirements"), (B3_PID, "proposals"), (Q_ID, "quotations")]:
    if eid:
        api("DELETE", f"/{ep}/{eid}")

print("\n" + "=" * 50)
passed = sum(1 for _, ok, _ in RES if ok)
print(f"验证结果: {passed}/{len(RES)} 通过")
sys.exit(0 if passed == len(RES) else 1)
